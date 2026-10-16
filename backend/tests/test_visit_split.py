"""A page that records several separate dated visits is summarized once per visit.

The case the reviewers raised (2026-10-15): one page holds several visits under the same heading,
each with its own date - a sheet of daily therapy or acupuncture notes - and they want each visit to
be its own entry in the report. Identification still makes the page one row; the summary call marks
each visit, `split_visits` cuts the body at the marks, and every visit is stored, re-drafted and
delivered as its own summary.

Synthetic data only: no model, no OCR. The model's reply is a stub carrying visit marks.
"""

import io
import time
import uuid

import docx as docxlib
import pytest
from sqlalchemy import select

from app.auth.password import MrrPasswordHelper
from app.db import get_sessionmaker
from app.models import Document, Job, ReviewRow, Summary, User
from app.services import bundles, jobs
from app.services import summarize_engine as se
from app.services.seed_catalog import constants_categories
from app.worker.tasks import summarize_document
from tests.conftest import unique_test_email

_NO_ISSUES = {"fixed_text": "", "issues": [], "ok": True}

_TWO_VISITS = (
    "**VISIT: 06/19/2026**\n"
    "**Body part being treated**: Low back. **Treatment provided**: Acupuncture.\n"
    "VISIT: 06/12/2026\n"
    "**Body part being treated**: Low back.\n"
    "**Treatment provided**: Cupping.\n"
)


def _row(**over):
    row = {
        "start": 19,
        "end": 19,
        "category": "5",
        "date": "06/19/2026",
        "injury_date": "-",
        "flag": "",
    }
    row.update(over)
    return row


def _clean_report(pages):
    listed = sorted({int(p) for p in pages})
    return {"pages": len(listed), "errored": [], "blank": []}


def _stub_model(monkeypatch, body, audit=None):
    """OCR and both model calls stubbed: the body call answers ``body``, the title call a header."""
    calls = []

    def fake_generate(model, system_msg, contents, temperature, max_output_tokens=None):
        calls.append(system_msg)
        if system_msg == se.TITLE_PROMPT:
            return "JANE ROE, L.AC. ACME ACUPUNCTURE. TREATMENT NOTE", False
        return body, False

    monkeypatch.setattr(
        se,
        "extract_pages_with_report",
        lambda path, pages, mark_pages=False, **_kw: ("OCR text", _clean_report(pages)),
    )
    monkeypatch.setattr(se, "_generate", fake_generate)
    monkeypatch.setattr(se, "verify_summary", audit or (lambda *a, **k: _NO_ISSUES))
    return calls


# ---------------------------------------------------------------------------------------------
# split_visits: cutting a marked body
# ---------------------------------------------------------------------------------------------


def test_a_marked_body_is_cut_into_its_visits_in_date_order():
    visits = se.split_visits(_TWO_VISITS, "06/19/2026")

    assert [date for date, _ in visits] == ["06/12/2026", "06/19/2026"]
    assert (
        visits[0][1] == "**Body part being treated**: Low back.\n**Treatment provided**: Cupping."
    )
    assert visits[1][1].startswith("**Body part being treated**: Low back. **Treatment provided**")


def test_decoration_around_a_mark_and_text_run_onto_its_line_are_read_past():
    body = (
        "`VISIT: 07/09/2026` - **Subjective**: better.\n"
        "- **VISIT: 07/24/2026**: **Subjective**: worse."
    )
    visits = se.split_visits(body, "07/09/2026")

    assert visits == [
        ("07/09/2026", "**Subjective**: better."),
        ("07/24/2026", "**Subjective**: worse."),
    ]


def test_text_before_the_first_mark_joins_the_first_visit_rather_than_being_lost():
    body = "Treating provider notes.\nVISIT: 06/12/2026\nFirst.\nVISIT: 06/19/2026\nSecond."
    visits = se.split_visits(body, "06/19/2026")

    assert visits[0] == ("06/12/2026", "Treating provider notes. First.")


def test_two_marks_naming_one_date_are_one_visit():
    """The date is what makes a visit separate. Two marks for one date are one visit, and when that
    leaves a single visit the row is not split at all."""
    assert se.split_visits("VISIT: 06/12/2026\nA\nVISIT: 6/12/26\nB", "06/12/2026") == []

    visits = se.split_visits("VISIT: 06/12/2026\nA\nVISIT: 06/19/2026\nB\nVISIT: 06/12/2026\nC")
    assert visits == [("06/12/2026", "A C"), ("06/19/2026", "B")]


@pytest.mark.parametrize(
    "body",
    [
        "**Treatment provided**: Acupuncture.",
        "VISIT: 06/12/2026\n**Treatment provided**: Acupuncture.",
        "VISIT: 06/12/2026\nA\nVISIT: 06/19/2026\n",
    ],
    ids=["no mark", "one mark", "a mark with no text"],
)
def test_fewer_than_two_visits_is_not_a_split(body):
    assert se.split_visits(body, "06/12/2026") == []


def test_a_mark_without_a_year_takes_the_year_nearest_the_rows_own_date():
    """Across a new year the nearest year is the previous one, not the row's."""
    visits = se.split_visits("VISIT: 12/28\nA\nVISIT: 1/4\nB", "01/10/2026")

    assert [date for date, _ in visits] == ["12/28/2025", "01/04/2026"]


def test_a_two_digit_year_is_written_with_four_digits():
    visits = se.split_visits("VISIT: 6/12/26\nA\nVISIT: 6/19/26\nB", "06/19/2026")

    assert [date for date, _ in visits] == ["06/12/2026", "06/19/2026"]


def test_a_mark_that_is_no_calendar_date_is_kept_as_written_and_sorts_last():
    """Never guessed at: the report shows it Undated, after the dated visits."""
    visits = se.split_visits("VISIT: 13/45/2026\nA\nVISIT: 06/12/2026\nB\nVISIT:\nC")

    assert visits == [("06/12/2026", "B"), ("13/45/2026", "A"), ("-", "C")]


def test_words_on_a_mark_that_carries_no_date_are_kept_as_the_visits_text():
    """Text the model put on a mark's line is never dropped, date or no date - losing a sentence of
    a summary is worse than a stray word in it."""
    visits = se.split_visits("VISIT: 06/12/2026\nA\nVISIT: **Subjective**: pain.\nB")

    assert visits == [("06/12/2026", "A"), ("-", "**Subjective**: pain. B")]


def test_a_lone_mark_is_removed_and_the_text_on_its_line_is_kept():
    stripped = se.strip_visit_lines("**VISIT: 06/12/2026** **Subjective**: pain.\nMore.")

    assert "VISIT" not in stripped
    assert stripped == "**Subjective**: pain.\nMore."


# ---------------------------------------------------------------------------------------------
# Which categories are told to mark visits
# ---------------------------------------------------------------------------------------------


def test_only_therapy_notes_and_diagnostic_studies_are_told_to_mark_visits():
    assert se.writes_visits("5")
    assert se.writes_visits("3")
    for category in ("1", "2", "9", "13", "14", "100", "an-id-nobody-made"):
        assert not se.writes_visits(category)
        assert se.VISIT_MARKER not in se.build_preamble(category), category


def test_the_instruction_and_the_split_read_one_set_for_every_catalog_category():
    """A category told to mark visits that nothing then splits would put the marks into a delivered
    summary; one split but never told would never split. Neither may happen for any category."""
    for category in constants_categories():
        told = se.VISIT_MARKER in se.build_preamble(category["id"])
        assert told == se.writes_visits(category["id"]), category["id"]


# ---------------------------------------------------------------------------------------------
# summarize_row: one output per visit
# ---------------------------------------------------------------------------------------------


def test_a_page_of_two_visits_summarizes_as_two_dated_summaries(monkeypatch):
    calls = _stub_model(monkeypatch, _TWO_VISITS)

    output = se.summarize_row("/x.pdf", _row(injury_date="08/20/2025"), prompt="P", verify=False)
    visits = se.visit_outputs(output)

    assert [visit["summaryDate"] for visit in visits] == ["06/12/2026", "06/19/2026"]
    # Each visit is a whole summary: its own body, the DOI the row states, one paragraph.
    assert visits[0]["summaryText"] == (
        "**DOI**: 08/20/2025. **Body part being treated**: Low back. "
        "**Treatment provided**: Cupping."
    )
    assert visits[1]["summaryText"].startswith("**DOI**: 08/20/2025. ")
    assert all("VISIT" not in visit["summaryText"] for visit in visits)
    # One heading for the page, shared by every visit.
    assert {visit["summaryTitle"] for visit in visits} == {
        "JANE ROE, L.AC. ACME ACUPUNCTURE. TREATMENT NOTE (Pages 19-19)"
    }
    # The top level is the first visit, so a caller reading one summary reads a correct one.
    assert output["summaryDate"] == "06/12/2026"
    assert output["summaryText"] == visits[0]["summaryText"]
    # And the instruction reached the body call.
    assert any(se.VISIT_MARKER in system for system in calls if system != se.TITLE_PROMPT)


def test_each_visit_is_audited_against_its_own_body_and_date(monkeypatch):
    """The audit reads the visit's date as the document date wherever that rule applies, so the
    split is exercised on a category that is handed one."""
    audits = []

    def audit(model, source, summary, title=None, document_date=None):
        audits.append((summary, document_date))
        issue = [{"type": "unsupported", "detail": "x"}]
        return {
            "fixed_text": summary.replace("Low back", "Lumbar"),
            "fixed_title": f"{title} (AUDITED {len(audits)})",
            "issues": issue,
            "ok": True,
        }

    monkeypatch.setattr(se, "_VISIT_LOG_CATEGORIES", frozenset({"1"}))
    _stub_model(monkeypatch, _TWO_VISITS, audit=audit)

    visits = se.visit_outputs(
        se.summarize_row("/x.pdf", _row(category="1"), prompt="P", verify=True)
    )

    assert [date for _, date in audits] == ["06/12/2026", "06/19/2026"]
    assert "Cupping" in audits[0][0]
    assert "Cupping" not in audits[1][0]
    assert all(visit["verified"] for visit in visits)
    assert "Lumbar" in visits[0]["verifiedText"]
    assert "Cupping" in visits[0]["verifiedText"]
    # The page's title correction is taken once, from the first visit's audit.
    assert visits[0]["verifiedTitle"] == visits[1]["verifiedTitle"]
    assert "(AUDITED 1)" in visits[1]["verifiedTitle"]


def test_a_single_marked_visit_keeps_the_rows_date_and_loses_the_mark(monkeypatch):
    _stub_model(monkeypatch, "VISIT: 06/12/2026\n**Treatment provided**: Acupuncture.")

    output = se.summarize_row("/x.pdf", _row(), prompt="P", verify=False)

    assert "visits" not in output
    assert se.visit_outputs(output) == [output]
    assert output["summaryDate"] == "06/19/2026"
    assert output["summaryText"] == "**Treatment provided**: Acupuncture."


def test_a_category_not_told_to_mark_visits_is_never_split(monkeypatch):
    calls = _stub_model(monkeypatch, _TWO_VISITS)

    output = se.summarize_row("/x.pdf", _row(category="1"), prompt="P", verify=False)

    assert "visits" not in output
    assert all(se.VISIT_MARKER not in system for system in calls)


# ---------------------------------------------------------------------------------------------
# The worker stores one summary per visit
# ---------------------------------------------------------------------------------------------


def _make_doc(page_count=3) -> str:
    with get_sessionmaker()() as session:
        user = User(
            email=unique_test_email(),
            name="Visits",
            password=MrrPasswordHelper().hash("Str0ng#pw1"),
            active=True,
        )
        session.add(user)
        session.flush()
        document = Document(
            id=str(uuid.uuid4()),
            user_id=user.id,
            original_filename="synthetic.pdf",
            stored_path="/nonexistent/synthetic.pdf",
            sha256="0" * 64,
            page_count=page_count,
        )
        session.add(document)
        session.commit()
        return document.id


def _seed_rows(doc_id, pages):
    with get_sessionmaker()() as session:
        for idx, page in enumerate(pages):
            session.add(
                ReviewRow(
                    document_id=doc_id,
                    idx=idx,
                    start=page,
                    end=page,
                    category="5",
                    title="Treatment Note",
                    date="-",
                    injury_date="-",
                    flag="-",
                    include=True,
                )
            )
        session.commit()


def _visit_output(page, dates):
    """A summarize_row stand-in's output for a row of ``len(dates)`` visits."""
    visits = [
        {
            "summaryTitle": f"T{page} (Pages {page}-{page})",
            "summaryDate": date,
            "summaryText": f"body {page} {date}",
            "manualCheck": "",
            "sourceText": "x",
        }
        for date in dates
    ]
    return visits[0] if len(visits) == 1 else {**visits[0], "visits": visits}


def _stored(doc_id):
    with get_sessionmaker()() as session:
        return [
            (s.idx, s.row_start, s.part, s.date)
            for s in session.scalars(
                select(Summary).where(Summary.document_id == doc_id).order_by(Summary.idx)
            ).all()
        ]


def _summarize(doc_id):
    with get_sessionmaker()() as session:
        job_id = jobs.create_job(session, doc_id, "summarize", model="m", prompt_version="1").id
    summarize_document(job_id)
    with get_sessionmaker()() as session:
        job = session.get(Job, job_id)
        assert job is not None
        assert job.state == "done"


def test_a_page_of_three_visits_is_stored_as_three_summaries_in_delivered_order(monkeypatch):
    """The split row finishes LAST, after the row below it already has its number, so placing it
    has to move that summary down - the order must still be row, then visit."""
    dates = {1: ["01/01/2026"], 2: ["02/01/2026", "02/08/2026", "02/15/2026"], 3: ["03/01/2026"]}

    def fake(pdf_path, row, model=None, prompt=None, standalone_studies=None, **_kw):
        time.sleep(0.05 if row["start"] == 2 else 0)
        return _visit_output(row["start"], dates[row["start"]])

    monkeypatch.setattr(se, "summarize_row", fake)
    doc_id = _make_doc()
    _seed_rows(doc_id, [1, 2, 3])

    _summarize(doc_id)

    assert _stored(doc_id) == [
        (0, 1, 0, "01/01/2026"),
        (1, 2, 0, "02/01/2026"),
        (2, 2, 1, "02/08/2026"),
        (3, 2, 2, "02/15/2026"),
        (4, 3, 0, "03/01/2026"),
    ]


def test_a_resume_keeps_every_visit_of_a_row_and_deletes_only_a_true_duplicate(monkeypatch):
    """Keyed by row alone, a resume kept one visit and deleted the others as its duplicates."""
    doc_id = _make_doc()
    _seed_rows(doc_id, [1, 2])
    with get_sessionmaker()() as session:
        job = Job(document_id=doc_id, kind="summarize", state="done", model="m", prompt_version="1")
        session.add(job)
        session.flush()
        for idx, part, text in ((7, 0, "visit A"), (8, 1, "visit B"), (9, 0, "duplicate of A")):
            session.add(
                Summary(
                    document_id=doc_id,
                    job_id=job.id,
                    idx=idx,
                    part=part,
                    title="T",
                    text=text,
                    row_start=1,
                    row_end=1,
                    row_category="5",
                )
            )
        session.commit()

    def fake(pdf_path, row, model=None, prompt=None, standalone_studies=None, **_kw):
        assert row["start"] == 2, "a row whose visits are all stored must not be summarized again"
        return _visit_output(2, ["05/01/2026"])

    monkeypatch.setattr(se, "summarize_row", fake)

    _summarize(doc_id)

    with get_sessionmaker()() as session:
        kept = session.scalars(
            select(Summary).where(Summary.document_id == doc_id).order_by(Summary.idx)
        ).all()
        assert [(s.idx, s.row_start, s.part, s.text) for s in kept] == [
            (0, 1, 0, "visit A"),
            (1, 1, 1, "visit B"),
            (2, 2, 0, "body 2 05/01/2026"),
        ]


# ---------------------------------------------------------------------------------------------
# Re-draft, export and the bundle
# ---------------------------------------------------------------------------------------------


@pytest.fixture
async def authed(client, seeded_user):
    email, password = seeded_user
    resp = await client.post("/api/auth/login", data={"username": email, "password": password})
    assert resp.status_code == 204
    return client


@pytest.fixture(autouse=True)
def _tmp_uploads(tmp_path, monkeypatch):
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "upload_folder", str(tmp_path))


async def _upload(client, pages=3) -> str:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=72, height=72)
    buffer = io.BytesIO()
    writer.write(buffer)
    resp = await client.post(
        "/api/documents", files={"pdf": ("scan.pdf", buffer.getvalue(), "application/pdf")}
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


def _seed_summaries(doc_id, specs):
    """``specs``: (idx, page, part, date) per summary, all on category 5 rows."""
    with get_sessionmaker()() as session:
        job = Job(document_id=doc_id, kind="summarize", state="done", model="m", prompt_version="1")
        session.add(job)
        session.flush()
        for idx, page, part, date in specs:
            session.add(
                Summary(
                    document_id=doc_id,
                    job_id=job.id,
                    idx=idx,
                    part=part,
                    title=f"TREATMENT NOTE {page}",
                    date=date,
                    text=f"old {page}.{part}",
                    row_start=page,
                    row_end=page,
                    row_category="5",
                )
            )
        session.commit()


def _redraft_as(monkeypatch, dates):
    import app.api.documents as documents_module

    def fake(_pdf_path, row, _model=None, prompt=None, standalone_studies=None):
        return _visit_output(row["start"], dates)

    monkeypatch.setattr(documents_module, "summarize_row", fake)


async def test_a_redraft_that_finds_more_visits_adds_summaries_and_moves_the_rest_down(
    authed, monkeypatch
):
    doc_id = await _upload(authed)
    _seed_summaries(doc_id, [(0, 1, 0, "01/01/2026"), (1, 2, 0, "02/01/2026")])
    _redraft_as(monkeypatch, ["01/01/2026", "01/08/2026", "01/15/2026"])

    resp = await authed.post(f"/api/documents/{doc_id}/summaries/0/resummarize")

    assert resp.status_code == 200, resp.text
    assert resp.json()["summaryDate"] == "01/01/2026"
    assert _stored(doc_id) == [
        (0, 1, 0, "01/01/2026"),
        (1, 1, 1, "01/08/2026"),
        (2, 1, 2, "01/15/2026"),
        (3, 2, 0, "02/01/2026"),
    ]
    listed = (await authed.get(f"/api/documents/{doc_id}/summaries")).json()
    assert [item["summaryDate"] for item in listed] == [
        "01/01/2026",
        "01/08/2026",
        "01/15/2026",
        "02/01/2026",
    ]


async def test_a_redraft_that_finds_fewer_visits_removes_the_extra_and_moves_the_rest_up(
    authed, monkeypatch
):
    doc_id = await _upload(authed)
    _seed_summaries(
        doc_id,
        [(0, 1, 0, "01/01/2026"), (1, 1, 1, "01/08/2026"), (2, 2, 0, "02/01/2026")],
    )
    _redraft_as(monkeypatch, ["01/03/2026"])

    resp = await authed.post(f"/api/documents/{doc_id}/summaries/1/resummarize")

    assert resp.status_code == 200, resp.text
    assert _stored(doc_id) == [(0, 1, 0, "01/03/2026"), (1, 2, 0, "02/01/2026")]


async def test_each_visit_is_its_own_entry_in_the_word_export_in_date_order(authed):
    doc_id = await _upload(authed)
    _seed_summaries(
        doc_id,
        [(0, 1, 0, "06/12/2026"), (1, 1, 1, "06/19/2026"), (2, 2, 0, "06/15/2026")],
    )
    resp = await authed.post(
        f"/api/documents/{doc_id}/export",
        json={"patientName": "Pat", "patientdob": "", "QMEorAME": "", "lawfirm": ""},
    )
    assert resp.status_code == 200, resp.text
    built = docxlib.Document(io.BytesIO((await authed.get(resp.json()["url"])).content))
    dates = [
        row.cells[0].text.strip()
        for table in built.tables
        for row in table.rows
        if row.cells[0].text.strip()[:2].isdigit()
    ]

    # Two visits of one page and a page between them in time, each its own dated entry.
    assert dates == ["06/12/26", "06/15/26", "06/19/26"]


async def test_the_bundle_cover_lists_a_page_of_visits_once_with_its_first_visits_date(authed):
    from app.api.documents import _delivered_entries

    doc_id = await _upload(authed)
    _seed_summaries(doc_id, [(0, 1, 0, "06/12/2026"), (1, 1, 1, "06/19/2026")])

    with get_sessionmaker()() as session:
        document = session.get(Document, doc_id)
        assert document is not None
        entries = _delivered_entries(session, document)

    assert list(entries) == [1]
    assert entries[1][0] == "06/12/2026"


def test_the_bundle_summarize_export_lists_each_visit(monkeypatch):
    monkeypatch.setattr(
        se,
        "summarize_row",
        lambda pdf_path, row, model=None, prompt=None, verify=None: _visit_output(
            row["start"], ["06/12/2026", "06/19/2026"]
        ),
    )

    entries = bundles.bundle_summary_entries("/x.pdf", [_row(start=1, end=1)])

    assert [entry["summaryDate"] for entry in entries] == ["06/12/2026", "06/19/2026"]
